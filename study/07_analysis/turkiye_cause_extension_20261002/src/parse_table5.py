from pathlib import Path
import json, warnings
import pandas as pd
import numpy as np
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
warnings.filterwarnings('ignore',message='Workbook contains no default style')
cols={5:('Urban','MV'),6:('Urban','LV'),7:('Urban','All'),8:('Suburban','MV'),9:('Suburban','LV'),10:('Suburban','All'),11:('Rural','MV'),12:('Rural','LV'),13:('Rural','All'),14:('All','All')}
records=[]
for file in sorted((ROOT/'raw').glob('sedas_table5_202*.xlsx')):
    year,month=map(int,file.stem.split('_')[-2:])
    w=load_workbook(file,read_only=True,data_only=True).active
    sections={};sec=None
    for row in w.iter_rows(values_only=True):
        label=str(row[2] or '').strip()
        if len(label)>1 and label[1]==')' and label[0] in 'ABCDEF':sec=label[0]
        if label.startswith('Genel Toplam') or label.startswith('Kullanıcı Sayıları'):
            if sec in sections: raise ValueError(f'Duplicate section {file} {sec}')
            sections[sec]=row
    assert set(sections)==set('ABCDEF'),(file,sections.keys())
    for c,(group,voltage) in cols.items():
        for notified,hs,fs in [('Unannounced','A','C'),('Announced','B','D')]:
            rec=dict(year=year,month=month,settlement=group,voltage=voltage,notification=notified,
                     users=sections['F'][c-1],annualized_scope='monthly',
                     frequency=sections[fs][c-1],hours=sections[hs][c-1]/60,
                     source_file=file.name,source_sheet=w.title)
            records.append(rec)
d=pd.DataFrame(records)
assert d.groupby(['year','settlement','voltage','notification']).size().eq(12).all()
assert d.groupby(['year','settlement','voltage','notification']).users.nunique().eq(1).all()
d['conditional_duration_h']=d.hours/d.frequency
d.to_csv(ROOT/'data/published_monthly_settlement_metrics.csv',index=False)
keys=['year','settlement','voltage','notification']
a=d.groupby(keys,as_index=False).agg(frequency=('frequency','sum'),hours=('hours','sum'),users=('users','first'))
a['conditional_duration_h']=a.hours/a.frequency
a.to_csv(ROOT/'data/published_annual_settlement_metrics.csv',index=False)
parts=[]
for keys,b in a[a.settlement.isin(['Urban','Rural'])].groupby(['year','voltage','notification']):
    b=b.set_index('settlement');u=b.loc['Urban'];r=b.loc['Rural']
    f=(r.frequency-u.frequency)*(r.conditional_duration_h+u.conditional_duration_h)/2
    t=(r.conditional_duration_h-u.conditional_duration_h)*(r.frequency+u.frequency)/2
    assert np.isclose(f+t,r.hours-u.hours)
    parts.append(dict(year=keys[0],voltage=keys[1],notification=keys[2],hour_gap=r.hours-u.hours,
        frequency_component_h=f,duration_component_h=t,frequency_component_share=f/(r.hours-u.hours),
        frequency_ratio=r.frequency/u.frequency,burden_ratio=r.hours/u.hours,
        duration_ratio=r.conditional_duration_h/u.conditional_duration_h))
pd.DataFrame(parts).to_csv(ROOT/'data/rural_urban_gap_partition.csv',index=False)
print(a[(a.voltage=='All')&(a.notification=='Unannounced')].to_string(index=False))
print(pd.DataFrame(parts).query("voltage == 'All' and notification == 'Unannounced'").to_string(index=False))
