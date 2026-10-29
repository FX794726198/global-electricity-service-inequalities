from pathlib import Path
import json, warnings
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
warnings.filterwarnings('ignore',message='Workbook contains no default style')
def tr_number(v):
    if v is None or v=='': return float('nan')
    if v=='#MULTIVALUE': return float('nan')
    if isinstance(v,str):
        return float(v.replace('.','').replace(',','.')) if ',' in v else float(v)
    return float(v)

def count_number(v):
    if v is None or v=='': return float('nan')
    if isinstance(v,str): return float(v.replace('.','').replace(',','.'))
    return float(v)

for year in [2024,2025]:
    path=ROOT/'raw'/f'sedas_tablo1_-2_{year}_-2.xlsx'
    w=load_workbook(path,read_only=True,data_only=True)
    rows=w.active.iter_rows(values_only=True)
    cols=next(rows)
    d=pd.DataFrame(rows,columns=cols)
    # Preserve source text: default read_excel inference mistakes Turkish thousands
    # separators for decimal points in columns containing only integer counts.
    d.to_parquet(ROOT/'data'/f'sedas_{year}_native.parquet',index=False)
    for c in cols[14:]:
        d[c]=d[c].map(count_number if c.startswith('EtkilenenKullanici') else tr_number)
    for c in cols[12:14]: d[c]=pd.to_datetime(d[c],errors='coerce')
    d['timestamp_duration_h']=(d[cols[13]]-d[cols[12]]).dt.total_seconds()/3600
    d['year']=year
    d['month']=d[cols[12]].dt.month
    d['source_file']=path.name
    d.to_parquet(ROOT/'data'/f'sedas_{year}_numeric.parquet',index=False)
    r=dict(year=year,rows=len(d),incident_ids=d.KesintininKodu1.nunique(),row_id_unique=d.Id.nunique(),
           provinces=d.Il3A.unique().tolist(),districts=d[['Il3A','Ilce3B']].drop_duplicates().shape[0],
           min_start=str(d[cols[12]].min()),max_start=str(d[cols[12]].max()),
           monthly_rows=d.groupby('month').size().to_dict(),
           max_duration_error_h=float((d[cols[14]]-d.timestamp_duration_h).abs().max()),
           negative_duration_rows=int((d[cols[14]]<0).sum()),missing_numeric_cells=int(d[list(cols[14:])].isna().sum().sum()))
    (ROOT/'metadata'/f'native_profile_{year}.json').write_text(json.dumps(r,indent=2,ensure_ascii=False))
    print(json.dumps(r,ensure_ascii=False),flush=True)
